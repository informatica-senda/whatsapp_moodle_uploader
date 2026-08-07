import csv
import json
import os
import re
import threading
import time
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

import requests
import psycopg
from openpyxl import load_workbook

from course_catalog import COURSE_MAP

APP_TITLE = "Senda - Moodle → PostgreSQL → WhatsApp"
ACCESS_URL_DEFAULT = "https://campusvirtual.sendagestion.com"
PLATFORM_DEFAULT = "moodle_senda"

REQUIRED_COLUMNS = [
    "username", "password", "firstname", "lastname", "email", "phone1",
    "course1", "profile_field_dni", "profile_field_empresa",
    "profile_field_inicio", "profile_field_fin"
]

DB_COLUMNS = [
    "external_id", "phone", "username", "password", "firstname", "lastname",
    "email", "dni", "course_name", "start_date", "end_date", "access_url",
    "platform"
]


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value):
    value = clean_text(value)
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha no válida: {value}")


def format_date_es(value):
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    return clean_text(value)


def normalize_phone(value):
    """Normaliza para BBDD (número nacional si es España) y API (E.164 sin +)."""
    digits = re.sub(r"\D", "", clean_text(value))
    if not digits:
        return "", ""

    # El flujo n8n actual guarda teléfonos españoles sin 34 en course_access.
    if digits.startswith("0034"):
        digits = digits[4:]
    elif digits.startswith("34") and len(digits) == 11:
        digits = digits[2:]

    db_phone = digits
    api_phone = digits if digits.startswith("34") and len(digits) > 9 else "34" + digits
    return db_phone, api_phone


def read_csv_file(path):
    last_error = None
    for enc in ("utf-8-sig", "cp1252", "latin1"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                # Los CSV de Moodle del usuario suelen venir separados por ;
                delimiter = ";" if sample.count(";") >= sample.count(",") else ","
                reader = csv.DictReader(f, delimiter=delimiter)
                return [dict(r) for r in reader]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise last_error or ValueError("No se pudo leer el CSV")


def read_xlsx_file(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [clean_text(v) for v in next(rows)]
    data = []
    for row in rows:
        if not any(v is not None and clean_text(v) for v in row):
            continue
        data.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
    return data


def read_input_file(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return read_csv_file(path)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx_file(path)
    raise ValueError("Formato no soportado. Usa CSV o XLSX.")


def standardize_course(short_name):
    key = clean_text(short_name)
    if key in COURSE_MAP:
        return COURSE_MAP[key]["name"], COURSE_MAP[key]["hours"]
    return "", ""


def prepare_rows(raw_rows, course_override=None, hours_override=None):
    if not raw_rows:
        return [], ["El fichero está vacío."]

    missing = [c for c in REQUIRED_COLUMNS if c not in raw_rows[0]]
    if missing:
        return [], ["Faltan columnas obligatorias: " + ", ".join(missing)]

    prepared = []
    errors = []
    for i, row in enumerate(raw_rows, start=2):
        try:
            username = clean_text(row.get("username"))
            password = clean_text(row.get("password"))
            firstname = clean_text(row.get("firstname"))
            lastname = clean_text(row.get("lastname"))
            email = clean_text(row.get("email"))
            dni = clean_text(row.get("profile_field_dni")) or username.upper()
            company = clean_text(row.get("profile_field_empresa"))
            course_short = clean_text(row.get("course1"))
            course_name, default_hours = standardize_course(course_short)
            if course_override:
                requested_name = clean_text(course_override)
                if not course_name or requested_name != course_name:
                    raise ValueError(
                        f"El nombre del curso debe configurarse en course_catalog.py: {course_short}"
                    )
            hours = clean_text(hours_override) if hours_override else default_hours
            start_date = parse_date(row.get("profile_field_inicio"))
            end_date = parse_date(row.get("profile_field_fin"))
            db_phone, api_phone = normalize_phone(row.get("phone1"))

            row_errors = []
            if not username: row_errors.append("username vacío")
            if not password: row_errors.append("password vacío")
            if not firstname: row_errors.append("firstname vacío")
            if not db_phone: row_errors.append("teléfono vacío")
            elif len(db_phone) != 9: row_errors.append(f"teléfono {db_phone} no tiene 9 dígitos")
            if not course_name: row_errors.append(f"curso no estandarizado: {course_short}")
            if email and not re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", email):
                row_errors.append(f"email sospechoso: {email}")

            prepared.append({
                "source_row": i,
                "external_id": username,  # criterio usado actualmente: DNI/username
                "phone": db_phone,
                "api_phone": api_phone,
                "username": username,
                "password": password,
                "firstname": firstname,
                "lastname": lastname,
                "email": email,
                "dni": dni,
                "company": company,
                "course_short": course_short,
                "course_name": course_name,
                "hours": hours,
                "start_date": start_date,
                "end_date": end_date,
                "access_url": ACCESS_URL_DEFAULT,
                "platform": PLATFORM_DEFAULT,
                "errors": row_errors,
                "status_db": "Pendiente",
                "status_wa": "Pendiente",
            })
            if row_errors:
                errors.append(f"Fila {i}: " + "; ".join(row_errors))
        except Exception as exc:
            errors.append(f"Fila {i}: {exc}")
    return prepared, errors


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1400x820")
        self.minsize(1100, 650)
        self.rows = []
        self.file_path = tk.StringVar()
        self.pg_url = tk.StringVar(value=os.getenv("SENDA_POSTGRES_URL", ""))
        self.wa_token = tk.StringVar(value=os.getenv("SENDA_WHATSAPP_TOKEN", ""))
        self.phone_number_id = tk.StringVar(value=os.getenv("SENDA_WHATSAPP_PHONE_ID", "1012284438640880"))
        self.template_name = tk.StringVar(value="confirmacion_cursos")
        self.template_language = tk.StringVar(value="es")
        self.course_override = tk.StringVar()
        self.hours_override = tk.StringVar()
        self.status_text = tk.StringVar(value="Selecciona un CSV/XLSX de Moodle.")
        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")

        ttk.Label(top, text="Archivo Moodle:").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.file_path, width=90).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(top, text="Seleccionar…", command=self.pick_file).grid(row=0, column=2)
        ttk.Button(top, text="Cargar y validar", command=self.load_file).grid(row=0, column=3, padx=(6,0))

        ttk.Label(top, text="PostgreSQL URL externa:").grid(row=1, column=0, sticky="w", pady=(8,0))
        ttk.Entry(top, textvariable=self.pg_url, show="•").grid(row=1, column=1, columnspan=3, sticky="ew", padx=6, pady=(8,0))

        ttk.Label(top, text="Token WhatsApp:").grid(row=2, column=0, sticky="w", pady=(8,0))
        ttk.Entry(top, textvariable=self.wa_token, show="•").grid(row=2, column=1, sticky="ew", padx=6, pady=(8,0))
        ttk.Label(top, text="Phone Number ID:").grid(row=2, column=2, sticky="e", padx=(10,0), pady=(8,0))
        ttk.Entry(top, textvariable=self.phone_number_id, width=22).grid(row=2, column=3, sticky="w", pady=(8,0))

        options = ttk.Frame(self, padding=(10,0,10,8))
        options.pack(fill="x")
        ttk.Label(options, text="Curso estandarizado (opcional):").grid(row=0, column=0, sticky="w")
        ttk.Entry(options, textvariable=self.course_override, width=50).grid(row=0, column=1, padx=6)
        ttk.Label(options, text="Horas (opcional):").grid(row=0, column=2, sticky="e")
        ttk.Entry(options, textvariable=self.hours_override, width=8).grid(row=0, column=3, padx=6)
        ttk.Label(options, text="Plantilla:").grid(row=0, column=4, sticky="e")
        ttk.Entry(options, textvariable=self.template_name, width=22).grid(row=0, column=5, padx=6)
        ttk.Label(options, text="Idioma:").grid(row=0, column=6, sticky="e")
        ttk.Entry(options, textvariable=self.template_language, width=6).grid(row=0, column=7, padx=6)

        actions = ttk.Frame(self, padding=(10,0,10,8))
        actions.pack(fill="x")
        ttk.Button(actions, text="1. Probar conexión PostgreSQL", command=lambda: self.run_thread(self.test_db)).pack(side="left")
        ttk.Button(actions, text="2. Subir alumnos a PostgreSQL", command=lambda: self.run_thread(self.upload_db)).pack(side="left", padx=6)
        ttk.Button(actions, text="3. Enviar WhatsApp masivo", command=lambda: self.run_thread(self.send_whatsapp)).pack(side="left", padx=6)
        ttk.Button(actions, text="Subir + enviar", command=lambda: self.run_thread(self.upload_and_send)).pack(side="left", padx=6)
        ttk.Button(actions, text="Exportar log CSV", command=self.export_log).pack(side="right")

        columns = ("row", "name", "phone", "course", "company", "start", "end", "hours", "db", "wa", "issues")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", height=26)
        headings = {
            "row":"Fila", "name":"Alumno", "phone":"Teléfono", "course":"Curso estandarizado",
            "company":"Empresa", "start":"Inicio", "end":"Fin", "hours":"Horas",
            "db":"PostgreSQL", "wa":"WhatsApp", "issues":"Avisos"
        }
        widths = {"row":50,"name":190,"phone":105,"course":280,"company":190,"start":90,"end":90,"hours":60,"db":100,"wa":100,"issues":300}
        for c in columns:
            self.tree.heading(c, text=headings[c])
            self.tree.column(c, width=widths[c], anchor="w")
        yscroll = ttk.Scrollbar(self, orient="vertical", command=self.tree.yview)
        xscroll = ttk.Scrollbar(self, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.pack(fill="both", expand=True, padx=(10,0))
        yscroll.place(relx=1.0, rely=0.28, relheight=0.65, anchor="ne")
        xscroll.pack(fill="x", padx=10)

        status = ttk.Label(self, textvariable=self.status_text, padding=10, relief="sunken", anchor="w")
        status.pack(fill="x", side="bottom")

        top.columnconfigure(1, weight=1)

    def pick_file(self):
        path = filedialog.askopenfilename(filetypes=[("Moodle CSV/Excel", "*.csv *.xlsx *.xlsm"), ("Todos", "*.*")])
        if path:
            self.file_path.set(path)
            # Al cambiar de archivo, volvemos a la estandarización automática del curso.
            self.course_override.set("")
            self.hours_override.set("")
            self.load_file()

    def load_file(self):
        try:
            raw = read_input_file(self.file_path.get())
            rows, errors = prepare_rows(raw, self.course_override.get(), self.hours_override.get())
            self.rows = rows
            # Si todo el fichero es del mismo curso, rellena los campos visibles automáticamente.
            courses = sorted({r["course_name"] for r in rows if r["course_name"]})
            hours = sorted({r["hours"] for r in rows if r["hours"]})
            if len(courses) == 1 and not self.course_override.get():
                self.course_override.set(courses[0])
            if len(hours) == 1 and not self.hours_override.get():
                self.hours_override.set(hours[0])
            self.refresh_tree()
            self.status_text.set(f"Cargados {len(rows)} alumnos. Avisos: {len(errors)}.")
            if errors:
                messagebox.showwarning("Validación", "Se han detectado avisos. Revísalos en la última columna antes de enviar.")
        except Exception as exc:
            messagebox.showerror("Error al cargar", str(exc))

    def refresh_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for idx, r in enumerate(self.rows):
            full_name = clean_text(f'{r["firstname"]} {r["lastname"]}')
            self.tree.insert("", "end", iid=str(idx), values=(
                r["source_row"], full_name, r["phone"], r["course_name"], r["company"],
                format_date_es(r["start_date"]), format_date_es(r["end_date"]), r["hours"],
                r["status_db"], r["status_wa"], "; ".join(r["errors"])
            ))

    def run_thread(self, fn):
        threading.Thread(target=self._thread_wrapper, args=(fn,), daemon=True).start()

    def _thread_wrapper(self, fn):
        try:
            fn()
        except Exception as exc:
            self.after(0, lambda: messagebox.showerror("Error", str(exc)))

    def test_db(self):
        url = self.pg_url.get().strip()
        if not url:
            raise ValueError("Pega la URL externa de PostgreSQL de EasyPanel.")
        self.status_text.set("Probando conexión PostgreSQL…")
        with psycopg.connect(url, connect_timeout=10) as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT 1")
                cur.fetchone()
        self.status_text.set("Conexión PostgreSQL correcta.")
        self.after(0, lambda: messagebox.showinfo("PostgreSQL", "Conexión correcta."))

    def valid_rows_for_action(self, require_hours=False):
        valid = []
        for r in self.rows:
            blocking = [e for e in r["errors"] if "email sospechoso" not in e]
            if require_hours and not r["hours"]:
                blocking.append("horas vacías")
            if not blocking:
                valid.append(r)
        return valid

    def upload_db(self):
        if not self.rows:
            raise ValueError("Carga primero el fichero Moodle.")
        url = self.pg_url.get().strip()
        if not url:
            raise ValueError("Pega la URL externa de PostgreSQL de EasyPanel.")

        valid = self.valid_rows_for_action(False)
        if not valid:
            raise ValueError("No hay filas válidas para subir.")

        self.status_text.set(f"Subiendo {len(valid)} alumnos a PostgreSQL…")
        inserted = skipped = failed = 0
        with psycopg.connect(url, connect_timeout=15) as conn:
            with conn.cursor() as cur:
                for r in valid:
                    try:
                        # Evita repetir exactamente la misma matrícula, pero permite que una persona tenga varios cursos.
                        cur.execute(
                            """
                            SELECT 1 FROM public.course_access
                            WHERE username = %s AND course_name = %s AND start_date = %s AND end_date = %s
                            LIMIT 1
                            """,
                            (r["username"], r["course_name"], r["start_date"], r["end_date"])
                        )
                        if cur.fetchone():
                            r["status_db"] = "Ya existía"
                            skipped += 1
                            continue

                        cur.execute(
                            """
                            INSERT INTO public.course_access (
                                external_id, phone, username, password, firstname, lastname, email, dni,
                                course_name, start_date, end_date, access_url, platform, created_at, updated_at
                            ) VALUES (
                                %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,NOW(),NOW()
                            )
                            """,
                            (r["external_id"], r["phone"], r["username"], r["password"], r["firstname"],
                             r["lastname"], r["email"], r["dni"], r["course_name"], r["start_date"],
                             r["end_date"], r["access_url"], r["platform"])
                        )
                        r["status_db"] = "Subido"
                        inserted += 1
                    except Exception as exc:
                        conn.rollback()
                        r["status_db"] = "ERROR"
                        r["errors"].append(f"DB: {exc}")
                        failed += 1
                    else:
                        conn.commit()
        self.after(0, self.refresh_tree)
        self.status_text.set(f"PostgreSQL: {inserted} subidos, {skipped} ya existían, {failed} errores.")

    def send_whatsapp(self):
        if not self.rows:
            raise ValueError("Carga primero el fichero Moodle.")
        token = self.wa_token.get().strip()
        phone_id = self.phone_number_id.get().strip()
        if not token or not phone_id:
            raise ValueError("Indica el token y el Phone Number ID de WhatsApp.")

        # Reaplica posibles overrides visibles sin tener que recargar el archivo.
        course_override = clean_text(self.course_override.get())
        hours_override = clean_text(self.hours_override.get())
        for r in self.rows:
            if course_override:
                r["course_name"] = course_override
            if hours_override:
                r["hours"] = hours_override

        valid = self.valid_rows_for_action(require_hours=True)
        if not valid:
            raise ValueError("No hay filas válidas. Revisa teléfono, curso y horas.")

        if not messagebox.askyesno("Confirmar envío", f"Se van a enviar {len(valid)} mensajes de WhatsApp. ¿Continuar?"):
            return

        url = f"https://graph.facebook.com/v25.0/{phone_id}/messages"
        headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        sent = failed = 0
        for pos, r in enumerate(valid, start=1):
            self.status_text.set(f"WhatsApp {pos}/{len(valid)}: {r['firstname']}…")
            payload = {
                "messaging_product": "whatsapp",
                "to": r["api_phone"],
                "type": "template",
                "template": {
                    "name": self.template_name.get().strip(),
                    "language": {"code": self.template_language.get().strip() or "es"},
                    "components": [{
                        "type": "body",
                        "parameters": [
                            {"type": "text", "text": r["firstname"].title()},
                            {"type": "text", "text": r["course_name"]},
                            {"type": "text", "text": r["company"]},
                            {"type": "text", "text": r["start_date"].strftime("%d/%m/%Y")},
                            {"type": "text", "text": r["end_date"].strftime("%d/%m/%Y")},
                            {"type": "text", "text": r["hours"]},
                            {"type": "text", "text": r["access_url"]},
                        ]
                    }]
                }
            }
            try:
                resp = requests.post(url, headers=headers, json=payload, timeout=30)
                if resp.ok:
                    data = resp.json()
                    msg_id = (data.get("messages") or [{}])[0].get("id", "")
                    r["status_wa"] = "Enviado"
                    if msg_id:
                        r["status_wa"] += f" ({msg_id[-8:]})"
                    sent += 1
                else:
                    r["status_wa"] = "ERROR"
                    r["errors"].append(f"WA {resp.status_code}: {resp.text[:300]}")
                    failed += 1
            except Exception as exc:
                r["status_wa"] = "ERROR"
                r["errors"].append(f"WA: {exc}")
                failed += 1
            self.after(0, self.refresh_tree)
            time.sleep(0.35)

        self.status_text.set(f"WhatsApp: {sent} enviados, {failed} errores.")
        self.after(0, lambda: messagebox.showinfo("WhatsApp", f"Proceso terminado.\nEnviados: {sent}\nErrores: {failed}"))

    def upload_and_send(self):
        self.upload_db()
        self.send_whatsapp()

    def export_log(self):
        if not self.rows:
            messagebox.showwarning("Sin datos", "No hay datos cargados.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV", "*.csv")], initialfile="log_envio_whatsapp.csv")
        if not path:
            return
        fields = ["source_row","username","firstname","lastname","phone","api_phone","course_name","company","start_date","end_date","hours","status_db","status_wa","errors"]
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, delimiter=";")
            w.writeheader()
            for r in self.rows:
                out = {k: r.get(k, "") for k in fields}
                out["start_date"] = format_date_es(r["start_date"])
                out["end_date"] = format_date_es(r["end_date"])
                out["errors"] = " | ".join(r["errors"])
                w.writerow(out)
        messagebox.showinfo("Log", f"Log guardado en:\n{path}")


if __name__ == "__main__":
    App().mainloop()
